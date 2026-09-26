from __future__ import annotations

import sys
from docx import Document


def paragraph_details(index, paragraph):
    print(f"P{index} style={paragraph.style.name!r} text={paragraph.text!r}")
    for run_index, run in enumerate(paragraph.runs):
        print(
            f"  R{run_index} text={run.text!r} bold={run.bold!r} italic={run.italic!r} "
            f"size={run.font.size.pt if run.font.size else None!r} font={run.font.name!r}"
        )


def cell_details(table_index, row_index, col_index, cell):
    print(f"T{table_index} R{row_index} C{col_index} text={cell.text!r}")
    for paragraph_index, paragraph in enumerate(cell.paragraphs):
        print(f"  CP{paragraph_index} style={paragraph.style.name!r} text={paragraph.text!r}")
        for run_index, run in enumerate(paragraph.runs):
            print(
                f"    R{run_index} text={run.text!r} bold={run.bold!r} "
                f"size={run.font.size.pt if run.font.size else None!r} font={run.font.name!r}"
            )


def main():
    document = Document(sys.argv[1])
    for start, end in [(20, 70), (115, 160)]:
        print(f"\nPARAGRAPHS {start}-{end}")
        for index in range(start, end + 1):
            paragraph_details(index, document.paragraphs[index])

    print("\nTARGET TABLES")
    for table_index in [0, 1, 2, 3, 4, 5]:
        table = document.tables[table_index]
        wanted_rows = range(len(table.rows))
        if table_index == 1:
            wanted_rows = range(0, 18)
        for row_index in wanted_rows:
            for col_index, cell in enumerate(table.rows[row_index].cells):
                cell_details(table_index, row_index, col_index, cell)


if __name__ == "__main__":
    main()
