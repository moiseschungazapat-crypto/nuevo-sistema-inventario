from __future__ import annotations

import shutil
import sys
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


BLUE = RGBColor(0x10, 0x36, 0x75)


def clear_paragraph(paragraph):
    """Remove paragraph contents while keeping its paragraph properties."""
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, *, name="Arial", size=10, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_paragraph(paragraph, text, *, size=10, bold=False, alignment=None, color=None):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if alignment is not None:
        paragraph.alignment = alignment
    return run


def set_labeled_paragraph(paragraph, parts, *, size=12):
    """parts is a list of (text, bold)."""
    clear_paragraph(paragraph)
    for text, bold in parts:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold)


def set_cell_text(cell, text, *, size=9, bold=False, alignment=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    set_paragraph(paragraph, text, size=size, bold=bold, alignment=alignment)
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)


def set_day_cell(cell, day, date):
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    day_run = paragraph.add_run(day)
    set_run_font(day_run, size=9.5, bold=True)
    day_run.add_break()
    date_run = paragraph.add_run(date)
    set_run_font(date_run, size=9)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)


def add_keep_together(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def patch_direction_textboxes(path: Path):
    with zipfile.ZipFile(path, "r") as archive:
        items = [(info, archive.read(info.filename)) for info in archive.infolist()]

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    patched = []
    for info, data in items:
        if info.filename == "word/document.xml":
            root = etree.fromstring(data)
            for textbox in root.xpath(".//w:txbxContent", namespaces=ns):
                texts = textbox.xpath(".//w:t", namespaces=ns)
                joined = "".join(node.text or "" for node in texts)
                if joined.startswith("DIRECCIÓN ZONAL"):
                    for node in texts:
                        if node.text and "_" in node.text:
                            node.text = "LIMA - CALLAO"
                            break
            data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
        patched.append((info, data))

    with NamedTemporaryFile(delete=False, suffix=".docx", dir=path.parent) as tmp:
        tmp_path = Path(tmp.name)
    try:
        with zipfile.ZipFile(tmp_path, "w") as archive:
            for info, data in patched:
                archive.writestr(info, data)
        shutil.move(str(tmp_path), str(path))
    finally:
        tmp_path.unlink(missing_ok=True)


def build(source: Path, output: Path, diagram: Path):
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(source)

    # Portada e identificación del estudiante.
    set_paragraph(doc.paragraphs[29], "CÓDIGO N° 02", size=14, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_labeled_paragraph(doc.paragraphs[50], [("CFP/UCP/ESCUELA: ", True), ("SENATI SEDE INDEPENDENCIA", False)])
    set_labeled_paragraph(doc.paragraphs[53], [("ESTUDIANTE: ", True), ("JACOB MOISES CHUNGA ZAPATA", False)])
    set_labeled_paragraph(
        doc.paragraphs[56],
        [("ID: ", True), ("001484105", False), ("    BLOQUE: ", True), ("NRC_47773", False)],
    )
    set_labeled_paragraph(doc.paragraphs[59], [("CARRERA: ", True), ("INGENIERÍA DE SOFTWARE CON IA", False)])
    set_labeled_paragraph(
        doc.paragraphs[62],
        [("INSTRUCTOR: ", True), ("GEORGE LEONARD TENORIO GONZALES", False)],
    )
    set_labeled_paragraph(
        doc.paragraphs[66],
        [("SEMESTRE: ", True), ("1ERO", False), ("    DEL: ", True), ("07/09/2026", False), ("    AL: ", True), ("19/09/2026", False)],
    )

    # Plan de rotaciones.
    rotations = doc.tables[0]
    set_cell_text(rotations.cell(3, 0), "LA LIGURIA S.A.\nÁREA DE SISTEMAS / DESARROLLO DE SOFTWARE", size=9, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(rotations.cell(3, 1), "07/09/2026", size=8, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(rotations.cell(3, 2), "19/09/2026", size=8, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(rotations.cell(3, 3), "2", size=10, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Plan específico de aprendizaje (PEA).
    pea_items = [
        "Realiza operaciones con las Librerías Pandas y Numpy",
        "Estudia el uso de las Librerías Scikit-learn y Pytorch",
        "Estudia el uso de las Librerías SciPy y Nltk",
        "Estudia el uso de las Librerías Tensorflow y Keras",
        "Realiza operaciones con las librerías Matplotlib y Seaborn",
        "Estudia los fundamentos de Inteligencia Artificial",
        "Realiza operaciones con algebra lineal, vectores y matrices",
        "Estudia los principios y variables estadísticas",
        "Realiza operaciones con la varianza y desviación estándar",
        "Crea programas con algoritmos de aprendizaje supervisado",
        "Crea programas con algoritmos de aprendizaje no supervisado",
        "Define la estructura y crea una red neuronal artificial",
        "Describe los principios de visión computacional y Machine Learning",
        "Explica el uso de la IA",
        "Utiliza recursividad en la programación",
        "Construye algoritmo de árbol de decisiones",
    ]
    pea = doc.tables[1]
    for index, description in enumerate(pea_items, start=1):
        row = index + 1
        pea.rows[row].height = Pt(17)
        pea.rows[row].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        set_cell_text(pea.cell(row, 0), f"{index:02d}", size=8.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(pea.cell(row, 1), description, size=7, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(pea.cell(row, 2), "X", size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(pea.cell(row, 3), "X" if index in (8, 14) else "", size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        for column in range(4, 7):
            set_cell_text(pea.cell(row, column), "", size=9, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Registro semanal de trabajos efectuados.
    week_1 = [
        ("LUNES", "07/09/2026", "Revisión del prototipo y detección de riesgos en el inicio de sesión y módulos incompletos.", "4"),
        ("MARTES", "08/09/2026", "Diseño del acceso con Supabase Auth, perfiles activos y separación de roles.", "4"),
        ("MIÉRCOLES", "09/09/2026", "Implementación de protección de páginas, cierre de sesión y recuperación de contraseña.", "4"),
        ("JUEVES", "10/09/2026", "Reorganización del menú lateral compartido y adaptación responsive de las secciones.", "4"),
        ("VIERNES", "11/09/2026", "Conexión del catálogo de productos con Supabase y validación segura de formularios.", "4"),
        ("SÁBADO", "12/09/2026", "Pruebas del acceso, revisión de errores y documentación del avance.", "4"),
    ]
    week_2 = [
        ("LUNES", "14/09/2026", "Diseño de inventario por sede, lotes, vencimientos y niveles mínimos.", "4"),
        ("MARTES", "15/09/2026", "Implementación de entradas, salidas, traslados y ajustes mediante funciones PostgreSQL.", "4"),
        ("MIÉRCOLES", "16/09/2026", "Validación de stock negativo, operaciones duplicadas y permisos por rol.", "4"),
        ("JUEVES", "17/09/2026", "Configuración del segundo paso OTP de 6 dígitos y políticas RLS.", "4"),
        ("VIERNES", "18/09/2026", "Integración de Realtime, auditoría, dashboard y reportes con datos reales.", "4"),
        ("SÁBADO", "19/09/2026", "Pruebas funcionales, corrección de incidencias y preparación del informe técnico.", "4"),
    ]
    for table, activities in ((doc.tables[2], week_1), (doc.tables[3], week_2)):
        for row, (day, date, activity, hours) in enumerate(activities, start=1):
            set_day_cell(table.cell(row, 0), day, date)
            set_cell_text(table.cell(row, 1), activity, size=9, alignment=WD_ALIGN_PARAGRAPH.LEFT)
            set_cell_text(table.cell(row, 2), hours, size=10, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(7, 1), "TOTAL", size=10, bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        set_cell_text(table.cell(7, 2), "24", size=10, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Informe técnico de la tarea significativa.
    significant = (
        "Fortalecimiento de la autenticación y del control transaccional del inventario en Supabase. "
        "Elegí esta tarea porque corrigió el acceso simulado del prototipo y permitió proteger los datos, "
        "controlar el stock por sede y lote y conservar la trazabilidad de cada movimiento. Se relaciona con "
        "las operaciones del PEA N.° 08 (principios y variables estadísticas aplicadas a indicadores de inventario) "
        "y N.° 14 (uso de la IA como apoyo para analizar, probar y corregir la solución)."
    )
    set_paragraph(doc.paragraphs[125], significant, size=9.3, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)

    steps = [
        "1. Revisar el prototipo e identificar riesgos en el login, las sesiones y la actualización directa del stock.",
        "2. Sustituir el acceso simulado por Supabase Auth con correo y contraseña.",
        "3. Crear perfiles activos y roles de administrador, operador y consulta.",
        "4. Aplicar RLS para que la base de datos controle las operaciones permitidas.",
        "5. Incorporar recuperación de contraseña y un segundo paso OTP de seis dígitos.",
        "6. Proteger las páginas y validar la identidad antes de consultar información.",
        "7. Modelar existencias por producto, sede y lote, con vencimiento y stock mínimo.",
        "8. Implementar entradas, salidas, traslados y ajustes mediante funciones PostgreSQL.",
        "9. Ejecutar cada movimiento en una transacción que actualiza stock e historial conjuntamente.",
        "10. Rechazar cantidades inválidas, stock negativo, lotes vencidos y reintentos duplicados.",
        "11. Registrar actor, sede, motivo, documento y saldos resultantes para auditoría.",
        "12. Conectar Dashboard y Reportes con información real y actualización mediante Realtime.",
        "13. Probar roles, recuperación, movimientos y visualización en distintos tamaños de pantalla.",
        "14. Documentar migraciones, activación y pruebas necesarias antes de producción.",
        "15. Corregir hallazgos y validar mensajes seguros sin exponer detalles internos.",
        "16. Preparar el despliegue en Vercel y la guía de configuración de Supabase.",
        "17. Confirmar que consulta solo lea y que los ajustes sean exclusivos del administrador.",
        "18. Registrar como pendientes las pruebas remotas de concurrencia y respaldo/recuperación.",
    ]
    for paragraph, step in zip(doc.paragraphs[127:145], steps):
        set_paragraph(paragraph, step, size=8.2, alignment=WD_ALIGN_PARAGRAPH.LEFT)

    doc.paragraphs[145].paragraph_format.page_break_before = True
    tools_1 = (
        "• Computadora con Windows, Edge/Chrome, Internet, Visual Studio Code, Git y GitHub.\n"
        "• HTML5, CSS3 y JavaScript modular para la interfaz y la lógica del cliente."
    )
    tools_2 = (
        "• Supabase Auth, PostgreSQL, RLS, Storage y Realtime para identidad, datos y sincronización.\n"
        "• Node.js, pnpm, PGlite, Playwright, Vercel y herramientas de IA como apoyo técnico."
    )
    set_paragraph(doc.paragraphs[146], tools_1, size=8.1, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    set_paragraph(doc.paragraphs[147], tools_2, size=8.1, alignment=WD_ALIGN_PARAGRAPH.LEFT)

    safety = (
        "Se mantuvo postura ergonómica y pausas activas. Las contraseñas, OTP y credenciales no se guardaron en el "
        "código ni se compartieron. Se usaron sesiones individuales, datos de prueba y un puesto ordenado y seguro."
    )
    heading = doc.paragraphs[148]
    original_heading = "Seguridad e higiene industrial/ambiental (ATS, Charla de cinco minutos: SST/SGA)"
    clear_paragraph(heading)
    head_run = heading.add_run(original_heading)
    set_run_font(head_run, size=11, bold=True)
    head_run.add_break()
    safety_run = heading.add_run(safety)
    set_run_font(safety_run, size=8.1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    results = (
        "Se reemplazó el acceso simulado por autenticación real con segundo factor, roles y RLS. Los movimientos actualizan "
        "stock e historial en una transacción y rechazan saldos negativos o solicitudes duplicadas."
    )
    results_heading = doc.paragraphs[149]
    original_results_heading = (
        "Resultados de la ejecución de la tarea/Recomendaciones (¿Se logró el objetivo que motivó la ejecución de la tarea? "
        "Qué recomendaciones sugiere para garantizar la operatividad del bien o servicio realizado)"
    )
    clear_paragraph(results_heading)
    head_run = results_heading.add_run(original_results_heading)
    set_run_font(head_run, size=11, bold=True)
    head_run.add_break()
    results_run = results_heading.add_run(results)
    set_run_font(results_run, size=8.1)
    results_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph(
        doc.paragraphs[150],
        "El sistema incorpora lotes, vencimientos, mínimos, auditoría y sincronización en tiempo real.",
        size=8.1,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    set_paragraph(
        doc.paragraphs[151],
        "Se recomienda probar la concurrencia, validar respaldos y restauración, mantener las migraciones versionadas y ejecutar pruebas de aceptación antes de producción.",
        size=8.1,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    # Esquema / diagrama del proceso.
    diagram_heading = doc.tables[4].cell(0, 0).paragraphs[0]
    add_keep_together(diagram_heading)
    diagram_heading.paragraph_format.page_break_before = True
    diagram_cell = doc.tables[4].cell(1, 0)
    paragraph = diagram_cell.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(diagram), width=Inches(6.15))
    diagram_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for extra in list(diagram_cell.paragraphs[1:]):
        diagram_cell._tc.remove(extra._p)

    # No se completan evaluación, calificación ni firmas: corresponden a terceros.
    doc.save(output)
    patch_direction_textboxes(output)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("usage: build_report.py SOURCE_DOCX OUTPUT_DOCX DIAGRAM_PNG")
    build(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]))
