from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


def main() -> None:
    source = Path(sys.argv[1])
    document = Document(source)
    print(f"sha256={hashlib.sha256(source.read_bytes()).hexdigest()}")
    print(f"sections={len(document.sections)} paragraphs={len(document.paragraphs)} tables={len(document.tables)} inline_shapes={len(document.inline_shapes)}")

    print("\nPARAGRAPHS")
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.replace("\n", "\\n").strip()
        if text:
            print(json.dumps({"index": index, "style": paragraph.style.name, "text": text}, ensure_ascii=False))

    print("\nTABLES")
    for table_index, table in enumerate(document.tables):
        print(f"TABLE {table_index}: rows={len(table.rows)} cols={len(table.columns)}")
        for row_index, row in enumerate(table.rows):
            cells = []
            for cell in row.cells:
                cells.append(" | ".join(p.text.strip().replace("\n", "\\n") for p in cell.paragraphs if p.text.strip()))
            print(json.dumps({"row": row_index, "cells": cells}, ensure_ascii=False))

    print("\nPACKAGE")
    package_inventory = []
    with zipfile.ZipFile(source) as archive:
        for info in archive.infolist():
            digest = hashlib.sha256(archive.read(info.filename)).hexdigest()
            package_inventory.append({"path": info.filename, "size": info.file_size, "sha256": digest})
            print(json.dumps({"path": info.filename, "size": info.file_size}, ensure_ascii=False))
        xml = etree.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        controls = xml.xpath(".//w:sdt", namespaces=ns)
        print(f"content_controls={len(controls)}")
        textboxes = xml.xpath(".//w:txbxContent", namespaces=ns)
        print(f"textboxes={len(textboxes)}")
        for index, textbox in enumerate(textboxes):
            text = "".join(textbox.xpath(".//w:t/text()", namespaces=ns)).strip()
            print(json.dumps({"textbox": index, "text": text}, ensure_ascii=False))

    if len(sys.argv) > 2:
        Path(sys.argv[2]).write_text(json.dumps(package_inventory, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
